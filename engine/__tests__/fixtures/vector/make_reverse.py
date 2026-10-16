"""Fixtures for v2.7 V22, the reverse check (engine/__tests__/vectorReverse.test.ts):
figures made from KNOWN DATA, so what is read from the drawing can be compared
with the numbers the figure was made from, not with another reading of it.

The data are written to 2 decimals (the "printed" data). Each truth file holds,
in PDF page points (y up):
  box     the plot area
  calib   two known values on each axis and where the generator put them
  series  per series: what it is drawn as, its colour, and its data

  python3 make_reverse.py            matplotlib (truth from ax.transData)
  uv run --with pymupdf --with openpyxl python3 make_reverse.py lo
                                     LibreOffice Calc (truth: see make_lo)

R's three ways are in make_reverse.R.
"""
import json
import os
import subprocess
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))


def data(rng, n, lo=1.0, hi=9.0):
    x = np.round(np.sort(rng.uniform(0.5, 9.5, n)), 2)
    y = np.round(rng.uniform(lo, hi, n), 2)
    return np.c_[x, y]


def curve(rng, n):
    x = np.round(np.linspace(0.5, 9.5, n), 2)
    a = rng.uniform(-1, 1, 3)
    y = np.round(5 + 2.5 * a[0] * np.sin(x / (1.2 + a[1])) + 0.25 * a[2] * x, 2)
    return np.c_[x, y]


def rgb(h):
    return [int(h[i:i + 2], 16) for i in (1, 3, 5)]


def write(name, truth):
    json.dump(truth, open(os.path.join(HERE, name + '.truth.json'), 'w'))


def mpl():
    import matplotlib
    matplotlib.use('pdf')
    import matplotlib.pyplot as plt
    matplotlib.rcParams['pdf.fonttype'] = 42

    def save(fig, ax, name, series):
        fig.canvas.draw()
        k = 72.0 / fig.dpi
        pt = lambda xy: (ax.transData.transform(np.asarray(xy, float)) * k).round(4).tolist()
        bb = ax.get_window_extent()
        (x0, x1), (y0, y1) = ax.get_xlim(), ax.get_ylim()
        xt = [v for v in ax.get_xticks() if x0 <= v <= x1]
        yt = [v for v in ax.get_yticks() if y0 <= v <= y1]
        calib = {
            'x': [[float(v), pt([v, y0])[0]] for v in (xt[0], xt[-1])],
            'y': [[float(v), pt([x0, v])[1]] for v in (yt[0], yt[-1])],
        }
        write(name, {'generator': 'matplotlib ' + matplotlib.__version__, 'digits': 2,
                     'box': [round(v * k, 4) for v in (bb.x0, bb.y0, bb.x1, bb.y1)], 'calib': calib, 'series': series})
        fig.savefig(os.path.join(HERE, name + '.pdf'))
        plt.close(fig)

    rng = np.random.default_rng(22001)
    # lines: two series in their own colours, plain lines
    fig, ax = plt.subplots(figsize=(4, 3))
    s = []
    for c in ('#1f77b4', '#d62728'):
        d = curve(rng, 30)
        ax.plot(d[:, 0], d[:, 1], color=c)
        s.append({'drawn': 'line', 'colour': rgb(c), 'data': d.tolist()})
    save(fig, ax, 'reverse-mpl-lines', s)

    # line and markers: the line through every data point, a marker on each
    fig, ax = plt.subplots(figsize=(4, 3))
    d = curve(rng, 15)
    ax.plot(d[:, 0], d[:, 1], color='#2ca02c', marker='o', markersize=4)
    save(fig, ax, 'reverse-mpl-linemarkers', [{'drawn': 'line+markers', 'colour': rgb('#2ca02c'), 'data': d.tolist()}])

    # scatter: two series, filled circles and squares
    fig, ax = plt.subplots(figsize=(4, 3))
    s = []
    for c, m in (('#9467bd', 'o'), ('#ff7f0e', 's')):
        d = data(rng, 20)
        ax.scatter(d[:, 0], d[:, 1], color=c, marker=m, s=25)
        s.append({'drawn': 'markers', 'colour': rgb(c), 'data': d.tolist()})
    save(fig, ax, 'reverse-mpl-scatter', s)

    # dense: a spectrum-like line whose corners are about 0.4 pt apart, closer
    # than Find lines' 2.5 px thinning at the app's render scale (design note,
    # "How many corners to keep": waits for David)
    fig, ax = plt.subplots(figsize=(4, 3))
    x = np.round(np.linspace(0.5, 9.5, 451), 2)
    y = np.round(5 + 3 * np.sin(x * 1.7) * np.exp(-0.15 * x), 2)
    ax.plot(x, y, color='#8c564b')
    save(fig, ax, 'reverse-mpl-dense', [{'drawn': 'line', 'colour': rgb('#8c564b'), 'data': np.c_[x, y].tolist()}])


def make_lo():
    """A Calc scatter chart, smoothing OFF, exported by LibreOffice itself.

    LibreOffice has no coordinate API, so the map from data to page is taken
    from the chart's own GRIDLINES: the axes are fixed (0 to 10, a gridline
    every 2), so the first and last gridline sit at known values. That map is
    the one thing read off the PDF, from grey 2-point gridlines the reader
    under test leaves out (V27: it is not in a series' colour)."""
    import pymupdf
    from openpyxl import Workbook
    from openpyxl.chart import Reference, ScatterChart, Series
    from openpyxl.chart.axis import ChartLines
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties

    rng = np.random.default_rng(22002)
    for name, marker in (('reverse-lo-lines', False), ('reverse-lo-linemarkers', True)):
        cols = ['1F77B4', 'D62728']
        ds = [curve(rng, 21) for _ in cols]
        wb = Workbook()
        sheet = wb.active
        ws = wb.create_sheet('data')  # the data on its own sheet: page 1 is the chart alone
        ws.cell(1, 1, 'x')
        for i in range(len(ds)):
            ws.cell(1, 2 + i, f'Sample {chr(65 + i)}')
        for j in range(len(ds[0])):
            ws.cell(2 + j, 1, float(ds[0][j, 0]))
            for i, d in enumerate(ds):
                ws.cell(2 + j, 2 + i, float(d[j, 1]))
        ch = ScatterChart()
        ch.scatterStyle = 'lineMarker'
        ch.width, ch.height = 16, 11
        xr = Reference(ws, min_col=1, min_row=2, max_row=len(ds[0]) + 1)
        for i, c in enumerate(cols):
            s = Series(Reference(ws, min_col=2 + i, min_row=1, max_row=len(ds[0]) + 1), xr, title_from_data=True)
            s.smooth = False
            s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=c, w=int(1.5 * 12700)))
            if marker:
                s.marker.symbol = 'circle'
                s.marker.size = 5
                s.marker.graphicalProperties = GraphicalProperties(solidFill=c, ln=LineProperties(solidFill=c))
            else:
                s.marker.symbol = 'none'
            ch.series.append(s)
        for a in (ch.x_axis, ch.y_axis):
            a.scaling.min, a.scaling.max, a.majorUnit = 0, 10, 2
            a.delete = False
        ch.x_axis.majorGridlines = ChartLines()
        sheet.add_chart(ch, 'A1')
        xlsx = os.path.join(HERE, name + '.xlsx')
        wb.save(xlsx)
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', HERE, xlsx], check=True, capture_output=True)
        os.remove(xlsx)
        page = pymupdf.open(os.path.join(HERE, name + '.pdf'))[0]
        H = page.rect.height
        vx, hy = set(), set()
        for d in page.get_drawings():
            c = d.get('color')
            if c is None or len(d['items']) != 1 or d['items'][0][0] != 'l' or max(c) - min(c) > 0.02 or c[0] < 0.5:
                continue
            p, q = d['items'][0][1], d['items'][0][2]
            if abs(p.x - q.x) < 0.01 and abs(p.y - q.y) > 50:
                vx.add(round(p.x, 4))
            if abs(p.y - q.y) < 0.01 and abs(p.x - q.x) > 50:
                hy.add(round(H - p.y, 4))  # PDF y is up
        vx, hy = sorted(vx), sorted(hy)
        assert len(vx) == 6 and len(hy) == 6, (name, vx, hy)
        write(name, {'generator': 'LibreOffice (soffice --convert-to pdf)', 'digits': 2,
                     'box': [vx[0], hy[0], vx[-1], hy[-1]],
                     'calib': {'x': [[0.0, vx[0]], [10.0, vx[-1]]], 'y': [[0.0, hy[0]], [10.0, hy[-1]]]},
                     'series': [{'drawn': 'line+markers' if marker else 'line', 'colour': rgb('#' + c), 'data': d.tolist()} for c, d in zip(cols, ds)]})


if __name__ == '__main__':
    if sys.argv[1:] == ['lo']:
        make_lo()
    else:
        mpl()
    print('written')
